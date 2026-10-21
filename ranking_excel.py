"""Admin-only Excel ranking template, validation, preview and atomic division replacement."""
import base64
import binascii
import hashlib
import io
import json
import zipfile
from portal import Error, require, rows

MAX_BYTES=512*1024
COLUMNS=('athlete_id','name','division','position')

def workbook_tools():
    try:
        import openpyxl
        import defusedxml  # Required parser protection; openpyxl detects and uses it.
    except ImportError:
        raise Error('Excel support is unavailable. Install requirements.txt and restart the server.',503)
    return openpyxl

def template(db,user):
    require(user,'admin')
    xlsx=workbook_tools()
    book=xlsx.Workbook();sheet=book.active;sheet.title='Rankings'
    sheet.append(COLUMNS)
    for row in db.execute('SELECT a.id,a.name,r.division,r.position FROM rankings r JOIN athletes a ON a.id=r.athlete_id ORDER BY r.division,r.position'):
        sheet.append(tuple(row))
    for row in sheet.iter_rows(min_row=2):
        for cell in row[1:3]:cell.data_type='s'
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for col,width in [('A',16),('B',32),('C',24),('D',16)]:sheet.column_dimensions[col].width=width
    for cell in sheet[1]:cell.font=xlsx.styles.Font(bold=True,color='FFFFFF');cell.fill=xlsx.styles.PatternFill('solid',fgColor='123C2B')
    reference=book.create_sheet('Fighters')
    reference.append(('athlete_id','name','division','current_position'))
    for row in db.execute('SELECT a.id,a.name,a.division,r.position FROM athletes a LEFT JOIN rankings r ON r.athlete_id=a.id AND r.division=a.division ORDER BY a.division,a.name'):
        reference.append(tuple(row))
    for row in reference.iter_rows(min_row=2):
        for cell in row[1:3]:cell.data_type='s'
    reference.freeze_panes='A2';reference.auto_filter.ref=reference.dimensions
    for col,width in [('A',16),('B',32),('C',24),('D',20)]:reference.column_dimensions[col].width=width
    instructions=book.create_sheet('Instructions')
    for line in [
        'RFC rankings upload — demo / sample data',
        'Edit the Rankings sheet only. Keep athlete_id, name, division, position headers.',
        'Use existing athlete IDs from the Fighters reference sheet. Name and division must match that athlete.',
        'Each represented division replaces that entire division ranking. Omitted divisions stay unchanged.',
        'Within each division: one athlete per row; unique positions 1, 2, 3 … without gaps.',
        'To add a currently unranked athlete, add its ID, exact name, division and desired position.',
        'Do not use formulas, macros, duplicate fighters, blank required cells or invented athlete IDs.',
        'Save as .xlsx. Maximum file size 512 KiB; maximum 2,000 ranking rows.',
        'Upload as app admin, review the preview and removals, then explicitly confirm publication.',
        'An empty file cannot erase rankings. Unlisted athletes in a represented division become unranked.',
        'This upload changes rankings only; it does not change fight records or create athletes.'
    ]:instructions.append([line])
    instructions.column_dimensions['A'].width=115
    buffer=io.BytesIO();book.save(buffer)
    return buffer.getvalue()

def fingerprint(db,divisions):
    current=[r for r in rows(db,'SELECT * FROM rankings ORDER BY division,position') if r['division'] in divisions]
    return hashlib.sha256(json.dumps(current,sort_keys=True).encode()).hexdigest()

def import_rankings(db,user,data):
    require(user,'admin')
    xlsx=workbook_tools()
    encoded=data.get('file','')
    if not isinstance(encoded,str) or len(encoded)>MAX_BYTES*4//3+8:
        raise Error('Upload an .xlsx workbook no larger than 512 KiB')
    try:raw=base64.b64decode(encoded,validate=True)
    except (binascii.Error,ValueError):raise Error('Invalid workbook encoding')
    if not raw or len(raw)>MAX_BYTES:raise Error('Upload an .xlsx workbook no larger than 512 KiB')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>100 or sum(i.file_size for i in archive.infolist())>10*1024*1024:
                raise Error('Workbook is too large when expanded')
            if any('vbaProject' in n for n in archive.namelist()):raise Error('Macro workbooks are not supported')
        book=xlsx.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    except Error:raise
    except Exception:raise Error('Cannot read this workbook. Use the downloaded .xlsx template.')
    clean=[];errors=[];seen_ids=set();seen_positions=set()
    athletes={a['id']:a for a in rows(db,'SELECT id,name,division FROM athletes')}
    try:
        if 'Rankings' not in book.sheetnames:raise Error('Workbook must contain the Rankings sheet')
        sheet=book['Rankings']
        if (sheet.max_row or 0)>2001 or (sheet.max_column or 0)>4:raise Error('Use only the four template columns and at most 2,000 ranking rows')
        iterator=sheet.iter_rows(max_col=4)
        header=next(iterator,())
        if tuple(c.value for c in header)!=COLUMNS:raise Error('Headers must be athlete_id, name, division, position, in that order')
        for index,cells in enumerate(iterator,2):
            values=[c.value for c in cells]
            if all(v is None for v in values):continue
            if any(c.data_type=='f' for c in cells):errors.append(f'Row {index}: formulas are not allowed');continue
            aid,name,division,position=values
            if isinstance(aid,bool) or not isinstance(aid,(int,float)) or aid!=int(aid) or int(aid) not in athletes:
                errors.append(f'Row {index}: athlete_id must identify an existing fighter');continue
            aid=int(aid);athlete=athletes[aid]
            if not isinstance(name,str) or name.strip()!=athlete['name'] or not isinstance(division,str) or division.strip()!=athlete['division']:
                errors.append(f'Row {index}: name and division must match athlete {aid}');continue
            division=division.strip()
            if isinstance(position,bool) or not isinstance(position,(int,float)) or position!=int(position) or not 1<=position<=2000:
                errors.append(f'Row {index}: position must be a whole number from 1 to 2000');continue
            position=int(position)
            if aid in seen_ids or (division,position) in seen_positions:
                errors.append(f'Row {index}: duplicate athlete or division position');continue
            seen_ids.add(aid);seen_positions.add((division,position))
            clean.append({'athlete_id':aid,'name':name.strip(),'division':division,'position':position})
    except Error:raise
    except Exception:raise Error('Invalid workbook contents. Start with the downloaded template.')
    finally:book.close()
    if errors:raise Error('; '.join(errors[:15]))
    if not clean:raise Error('The workbook has no ranking rows. Download a populated template first.')
    divisions=sorted({r['division'] for r in clean})
    for division in divisions:
        positions=sorted(r['position'] for r in clean if r['division']==division)
        if positions!=list(range(1,len(positions)+1)):raise Error(f'{division}: positions must start at 1 with no gaps')
    current=[r for r in rows(db,'SELECT * FROM rankings ORDER BY division,position') if r['division'] in divisions]
    previous={r['athlete_id']:r['position'] for r in current}
    preview=sorted([dict(r,previous_position=previous.get(r['athlete_id'])) for r in clean],key=lambda r:(r['division'],r['position']))
    removed=[dict(r,name=athletes[r['athlete_id']]['name']) for r in current if r['athlete_id'] not in seen_ids]
    revision=fingerprint(db,divisions)
    if data.get('preview',True) is not False:
        return {'preview':True,'rows':preview,'removed':removed,'divisions':divisions,'revision':revision}
    if data.get('revision')!=revision:raise Error('Rankings changed after preview. Upload again and review the new preview.',409)
    for division in divisions:db.execute('DELETE FROM rankings WHERE division=?',(division,))
    for row in clean:
        aid,position=row['athlete_id'],row['position']
        db.execute('INSERT INTO rankings(division,athlete_id,position,movement,updated_at) VALUES(?,?,?,?,CURRENT_TIMESTAMP)',(row['division'],aid,position,previous.get(aid,position)-position))
    return {'ok':True,'imported':len(clean),'removed':len(removed),'divisions':divisions}
